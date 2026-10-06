# -*- coding: utf-8 -*-
"""生成跨境智汇 AI 系统演示 PPT（5 页）"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- 主题色 ----------
PRIMARY = RGBColor(0x1F, 0x6F, 0xEB)    # 品牌蓝
DARK = RGBColor(0x1A, 0x23, 0x32)       # 深墨色
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF5, 0xF8, 0xFD)
GRAY = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x00, 0xB0, 0x9A)     # 青绿点缀

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

FONT = "Microsoft YaHei"


def add_rect(slide, l, t, w, h, color, line=None):
    """添加填充矩形"""
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def add_text(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.0):
    """添加文本框"""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return tb


def add_bullets(slide, l, t, w, h, items, size=16, color=DARK, spacing=1.15):
    """添加带项目符号的列表，items 为 (标题, 描述) 列表"""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (title, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = "▪ " + title
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY
        r1.font.name = FONT
        if desc:
            r2 = p.add_run()
            r2.text = "　" + desc
            r2.font.size = Pt(size - 1)
            r2.font.color.rgb = GRAY
            r2.font.name = FONT
    return tb


def header(slide, title, subtitle=None):
    """统一页头：顶部色带 + 标题"""
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.16), PRIMARY)
    add_text(slide, Inches(0.6), Inches(0.45), Inches(9), Inches(0.7),
             title, 26, DARK, bold=True)
    if subtitle:
        add_text(slide, Inches(0.6), Inches(1.05), Inches(12), Inches(0.4),
                 subtitle, 14, GRAY)


def footer(slide, page_no, total=5):
    """统一页脚"""
    add_text(slide, Inches(11.6), Inches(7.0), Inches(1.2), Inches(0.4),
             f"{page_no} / {total}", 12, GRAY, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.6), Inches(7.0), Inches(6), Inches(0.4),
             "跨境智汇 AI 系统 · 项目演示", 12, GRAY)


# ================= 第 1 页：封面 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, prs.slide_width, prs.slide_height, DARK)
add_rect(s, 0, Inches(5.9), prs.slide_width, Inches(0.06), ACCENT)
add_text(s, Inches(2.1), Inches(1.6), Inches(9), Inches(1.0),
         "跨境智汇 AI 系统", 54, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.2), Inches(2.7), Inches(11), Inches(0.6),
         "粤港甄选 · 企业级分布式智能化运营平台", 22, RGBColor(0xBF, 0xD1, 0xE8),
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.2), Inches(4.0), Inches(11), Inches(1.2),
         "微服务架构 · RAG 智能客服 · 全链路安全 · 双端前端",
         16, RGBColor(0x9A, 0xA7, 0xB8), align=PP_ALIGN.CENTER)
add_text(s, Inches(1.2), Inches(6.2), Inches(11), Inches(0.5),
         "基于 Java 25 · Spring Cloud Alibaba · DeepSeek 大模型", 14,
         RGBColor(0x8A, 0x97, 0xAA), align=PP_ALIGN.CENTER)

# ================= 第 2 页：项目概述与可行性 =================
s = prs.slides.add_slide(BLANK)
header(s, "项目概述与可行性分析", "面向跨境电商企业的智能化运营平台")
# 左侧：业务定位
add_rect(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.1), LIGHT_BG)
add_text(s, Inches(0.9), Inches(1.85), Inches(5.3), Inches(0.5),
         "业务定位", 18, DARK, bold=True)
add_bullets(s, Inches(0.9), Inches(2.4), Inches(5.3), Inches(4.1), [
    ("智能化运营", "覆盖商品、库存、订单、支付全链路"),
    ("AI 智能客服", "RAG 架构，融合企业知识库与大模型"),
    ("知识与培训", "企业知识沉淀 + 员工培训闭环"),
    ("组织与权限", "RBAC 多维组织架构管控"),
], size=15)
# 右侧：技术栈
add_rect(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.1), LIGHT_BG)
add_text(s, Inches(7.1), Inches(1.85), Inches(5.3), Inches(0.5),
         "技术栈", 18, DARK, bold=True)
add_bullets(s, Inches(7.1), Inches(2.4), Inches(5.3), Inches(4.1), [
    ("后端", "Java 25 · Spring Boot 4 · Spring Cloud Alibaba"),
    ("数据与中间件", "MySQL 8.4 · Redis · Nacos · Elasticsearch"),
    ("AI 能力", "DeepSeek / 豆包方舟（多供应商可切换）"),
    ("前端", "Vue 3 + Vite · admin / mall 双端"),
    ("工程与安全", "JaCoCo 覆盖率门禁 · JWT + HMAC + AES-GCM"),
], size=15)
footer(s, 2)

# ================= 第 3 页：系统架构 =================
s = prs.slides.add_slide(BLANK)
header(s, "系统架构", "微服务 + API 网关，服务间安全签名通信")
# 层标题
add_text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.4),
         "前端层 → 网关层 → 业务服务层 → 基础设施层", 15, GRAY)
# 网关
add_rect(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.85), PRIMARY)
add_text(s, Inches(1.2), Inches(2.28), Inches(11), Inches(0.5),
         "ygh-gateway (8080) · JWT 会话校验 · 路由转发 · 限流 · CORS", 16, WHITE, bold=True)
# 业务服务网格（当前实际运行的 14 个服务）
services = [
    ("gateway", "8080", "网关"), ("auth", "8081", "认证"), ("user", "8082", "用户"),
    ("system", "8083", "系统"), ("product", "8084", "商品"), ("knowledge", "8085", "知识库"),
    ("inventory", "8086", "库存"), ("order", "8087", "订单"), ("training", "8089", "培训"),
    ("ai", "8090", "AI 客服"), ("wallet", "8091", "钱包"), ("notification", "8092", "通知"),
    ("admin", "8093", "管理端"), ("search", "8094", "搜索"),
]
col, row = 0, 0
box_w, box_h = Inches(1.52), Inches(0.78)
gap_x, gap_y = Inches(0.10), Inches(0.12)
start_x, start_y = Inches(0.9), Inches(3.15)
for name, port, label in services:
    x = start_x + col * (box_w + gap_x)
    y = start_y + row * (box_h + gap_y)
    add_rect(s, x, y, box_w, box_h, LIGHT_BG)
    add_text(s, x, y + Inches(0.08), box_w, Inches(0.35), name, 12, PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x, y + Inches(0.42), box_w, Inches(0.3), f"{label} · {port}", 9, GRAY, align=PP_ALIGN.CENTER)
    col += 1
    if col >= 7:
        col = 0
        row += 1
# AI 客服高亮
add_rect(s, Inches(0.9), Inches(4.95), Inches(5.6), Inches(0.95), ACCENT)
add_text(s, Inches(1.1), Inches(5.07), Inches(5.2), Inches(0.75),
         "AI 智能客服（核心）\nRAG 检索 + 工具调用 + DeepSeek", 13, WHITE, bold=True)
# 基础设施
add_rect(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(0.7), DARK)
add_text(s, Inches(1.1), Inches(6.2), Inches(11.2), Inches(0.4),
         "基础设施：MySQL 8.4 · Redis 7 · Nacos · Elasticsearch · RocketMQ(可选)", 14, WHITE, bold=True)
footer(s, 3)

# ================= 第 4 页：核心功能（AI 客服） =================
s = prs.slides.add_slide(BLANK)
header(s, "核心亮点：RAG 智能客服", "多供应商大模型 + 企业知识库 + 实时业务数据")
# 步骤流程
steps = [
    ("① 用户提问", "识别问题中的 SKU / 订单号"),
    ("② 知识检索", "混合检索企业知识库片段"),
    ("③ 工具调用", "实时查询库存、订单数据"),
    ("④ 大模型生成", "DeepSeek 融合上下文作答"),
    ("⑤ 结果落地", "引用来源 + 会话持久化"),
]
width_per = Inches(11.3) / 5
for i, (t, d) in enumerate(steps):
    x = Inches(0.6) + i * width_per
    add_rect(s, x, Inches(1.7), width_per - Inches(0.15), Inches(1.5), LIGHT_BG)
    add_rect(s, x, Inches(1.7), width_per - Inches(0.15), Inches(0.12), PRIMARY)
    add_text(s, x + Inches(0.1), Inches(1.95), width_per - Inches(0.3), Inches(0.5),
             t, 15, DARK, bold=True)
    add_text(s, x + Inches(0.1), Inches(2.5), width_per - Inches(0.3), Inches(0.6),
             d, 11, GRAY)
# 特性
add_text(s, Inches(0.6), Inches(3.5), Inches(12), Inches(0.5),
         "关键特性", 18, DARK, bold=True)
add_bullets(s, Inches(0.6), Inches(4.05), Inches(12.1), Inches(2.6), [
    ("多供应商抽象", "DeepSeek / 豆包方舟动态切换，无需改业务代码"),
    ("容错降级", "检索或工具服务不可用时，AI 仍可基于通用知识作答"),
    ("安全合规", "对话内容分级可见性控制，敏感配置 AES-GCM 加密"),
    ("流式输出", "SSE 流式响应，分段返回提升体验"),
], size=15, spacing=1.2)
footer(s, 4)

# ================= 第 5 页：结束页 =================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, prs.slide_width, prs.slide_height, DARK)
add_rect(s, 0, Inches(3.4), prs.slide_width, Inches(0.06), ACCENT)
add_text(s, Inches(1.2), Inches(2.1), Inches(11), Inches(1.0),
         "谢谢观看", 50, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.2), Inches(3.7), Inches(11), Inches(0.6),
         "跨境智汇 AI 系统 —— 让企业运营更智能", 22, RGBColor(0xBF, 0xD1, 0xE8),
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.2), Inches(5.0), Inches(11), Inches(0.5),
         "微服务 · RAG AI 客服 · 全链路安全 · 双端前端", 16, RGBColor(0x9A, 0xA7, 0xB8),
         align=PP_ALIGN.CENTER)
footer(s, 5)

PRS_PATH = r"c:\Users\唐国几\IdeaProjects\yuegang-zhihui-ai1\跨境智汇AI系统项目演示.pptx"
prs.save(PRS_PATH)
print("PPT 已生成:", PRS_PATH)